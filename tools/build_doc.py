#!/usr/bin/env python3
"""Any note of the vault as an editable Word file and a PDF, diagrams included.

    python3 tools/build_doc.py "Knowledge/Decisions.md"               build/Decisions.docx and build/Decisions.pdf
    python3 tools/build_doc.py NOTE build/briefing --footer "Company · internal"

Reads the Markdown the notes use: frontmatter (left out), #, ## and ### headings, paragraphs, - bullets, numbered
items, tables, **bold**, a line <!-- page --> for a page break, a line ![caption](path) for a picture (the path
relative to the note; a missing picture leaves a frame with its caption), and ```mermaid blocks, drawn as figures (an
SVG in the PDF, a PNG in the Word file). Wikilinks become their text. A paragraph right after a diagram is its caption.
The colours come from "theme" in settings.json, the footer from "footer". Takes --vault PATH like vault.py.

Drawing and the PDF need Node and Playwright with its Chromium, as a Claude Code cloud session has them. Without
them the Word file is still built, without its diagrams.
"""
import html
import json
import re
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))
import vault as V  # noqa: E402

THEME = {"paper": "#FFFFFF", "ink": "#1F2328", "muted": "#59636E", "rule": "#D0D7DE", "soft": "#F6F8FA",
         "accent": "#0B5FAD", "alert": "#CF222E"}
BODY, MONO = "Archivo", "IBM Plex Mono"
EMPHASIS = r"(?<![*\w])\*([^*\s][^*]*?)\*(?![*\w])"  # *italics*, not **bold**
ID = re.compile(r"^[A-Z]{1,4}-?\d+$")  # an ID in a table's first column stays on one line


def rgb(hex_):
    return RGBColor.from_string(hex_.lstrip("#").upper())


def plain_links(text):
    def one(m):
        target, _, alias = m.group(1).partition("|")
        return alias or target.split("#")[0].split("/")[-1]
    return re.sub(r"\[\[([^\]]+)\]\]", one, text)


def parse(src):
    """The note as blocks: (kind, payload) with kind h1, h2, h3, p, ul, ol, start, table, page, image or diagram."""
    text = Path(src).read_text(encoding="utf-8")
    text = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S)
    text = re.sub(r"<!-- generated -->\n|<!-- /generated -->", "", text)
    text = re.sub(r"<!--(?! page -->).*?-->", "", text, flags=re.S)
    blocks, rows, fence = [], [], None
    for line in text.splitlines() + [""]:
        if fence is not None:
            if line.strip() == "```":
                if fence[0] == "mermaid":
                    blocks.append(("diagram", "\n".join(fence[1:])))
                fence = None
            else:
                fence.append(line)
            continue
        if line.startswith("```"):
            fence = [line[3:].strip()]
            continue
        s = plain_links(line.rstrip())
        if s.startswith("|"):
            cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", s.strip().strip("|"))]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
                rows.append([re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", c) for c in cells])
            continue
        if rows:
            blocks.append(("table", rows))
            rows = []
        if not s.strip():
            continue
        s = re.sub(r"(?<!!)\[([^\]]+)\]\([^)]+\)", r"\1", s)  # a Markdown link becomes its text
        img = re.fullmatch(r"!\[(.*?)\]\((.+?)\)", s.strip())
        if img:
            path = (Path(src).parent / img.group(2)).resolve()
            if not path.exists():
                print("  missing picture, a frame in its place: %s" % img.group(2))
            blocks.append(("image", {"caption": img.group(1), "path": path}))
        elif s.strip() == "<!-- page -->":
            blocks.append(("page", ""))
        elif s.startswith("# "):
            blocks.append(("h1", s[2:]))
        elif s.startswith("## "):
            blocks.append(("h2", s[3:]))
        elif s.startswith("### ") or s.startswith("#### "):
            blocks.append(("h3", s.lstrip("#").strip()))
        elif re.match(r"^\s*[-*] ", s):
            blocks.append(("ul", re.sub(r"^\s*[-*] ", "", s)))
        elif re.match(r"^\d+\. ", s):
            n = int(s.split(".")[0])
            if n > 1 and (not blocks or blocks[-1][0] != "ol"):
                blocks.append(("start", n))
            blocks.append(("ol", re.sub(r"^\d+\. ", "", s)))
        else:
            blocks.append(("p", s.strip()))
    return blocks


def why(e):
    """A short reason for a drawing or printing step that failed, such as Node missing."""
    text = str(getattr(e, "stderr", "") or e).strip()
    return text.splitlines()[-1][:120] if text else type(e).__name__


def png_width_px(path):
    with open(path, "rb") as f:
        return int.from_bytes(f.read(24)[16:20], "big")


def diagrams(blocks, folder, theme):
    """Draw every Mermaid block once; without Node and Playwright, leave them out with a note."""
    todo = [{"id": "fig%d" % (k + 1), "code": val} for k, val in enumerate(v for kind, v in blocks if kind == "diagram")]
    if not todo:
        return blocks
    folder.mkdir(parents=True, exist_ok=True)
    spec = folder / "diagrams.json"
    spec.write_text(json.dumps({"theme": theme, "items": todo}), encoding="utf-8")
    try:
        subprocess.run(["node", str(TOOLS / "render.js"), "mermaid", str(spec), str(folder)], check=True,
                       capture_output=True, text=True, timeout=300)
    except Exception as e:
        print("  diagrams not drawn (%s): they are left out" % why(e))
        return [(k, v) for k, v in blocks if k != "diagram"]
    files = iter({"svg": folder / (d["id"] + ".svg"), "png": folder / (d["id"] + ".png")} for d in todo)
    return [(kind, next(files)) if kind == "diagram" else (kind, val) for kind, val in blocks]


# ---------------- Word ----------------

def runs(par, text, size=10.5, color=None, bold=False, font=BODY):
    text = re.sub(r"`([^`]+)`", r"\1", re.sub(EMPHASIS, r"\1", text))  # Word gets the words, not the marks
    for k, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            r = par.add_run(chunk)
            r.font.name, r.font.size = font, Pt(size)
            if color is not None:
                r.font.color.rgb = color
            r.font.bold = bold or k % 2 == 1
            r._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        shd.set(qn(k), v)
    cell._element.get_or_add_tcPr().append(shd)


def word(blocks, out, theme, footer, title):
    ink, accent, muted = rgb(theme["ink"]), rgb(theme["accent"]), rgb(theme["muted"])
    doc = Document()
    props, now = doc.core_properties, datetime.combine(V.today(), datetime.now().time())
    props.title, props.author, props.last_modified_by, props.comments = title, "", "", ""
    props.created = props.modified = now
    props.revision = 1
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    runs(sec.footer.paragraphs[0], footer, 7, muted, font=MONO)
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = BODY, Pt(10.5)
    number = 0
    for kind, val in blocks:
        if kind == "start":
            number = val - 1
            continue
        number = number + 1 if kind == "ol" else 0
        if kind == "page":
            doc.add_page_break()
        elif kind == "image":
            if val["path"].exists():
                doc.add_picture(str(val["path"]), width=Cm(min(14.0, png_width_px(val["path"]) / 96 * 2.54)))
            p = doc.add_paragraph()
            runs(p, val["caption"] if val["path"].exists() else "[Picture to come] " + val["caption"], 8.5, muted)
        elif kind == "diagram":
            width_cm = min(17.0, png_width_px(val["png"]) / 3 / 96 * 2.54)  # drawn at three times, 96 px to the inch
            doc.add_picture(str(val["png"]), width=Cm(width_cm))
        elif kind == "table":
            cols = max(len(r) for r in val)
            t = doc.add_table(rows=len(val), cols=cols)
            t.style, t.alignment = "Table Grid", WD_TABLE_ALIGNMENT.LEFT
            for i, row in enumerate(val):
                for j in range(cols):
                    cell, txt = t.cell(i, j), row[j] if j < len(row) else ""
                    cell.paragraphs[0].paragraph_format.space_after = Pt(1)
                    if i == 0:
                        shade(cell, theme["ink"].lstrip("#"))
                        runs(cell.paragraphs[0], txt.upper(), 7, rgb(theme["paper"]), True, MONO)
                    else:
                        runs(cell.paragraphs[0], txt, 8.5, ink)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph(style="List Bullet") if kind == "ul" else doc.add_paragraph()
            if kind == "ol":
                p.paragraph_format.left_indent = Cm(0.6)
                val = "%d. %s" % (number, val)
            size, color, bold = {"h1": (20, ink, True), "h2": (13, accent, True), "h3": (11, ink, True)}.get(kind, (10.5, ink, False))
            if kind in ("h2", "h3"):
                p.paragraph_format.space_before = Pt(14 if kind == "h2" else 9)
                p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt({"h1": 8, "h2": 4, "h3": 3, "p": 6}.get(kind, 3))
            runs(p, val, size, color, bold)
    doc.save(out)


# ---------------- PDF through HTML ----------------

CSS = """
@font-face { font-family: Archivo; src: url('%(fonts)s/Archivo-Regular.ttf'); font-weight: 400; }
@font-face { font-family: Archivo; src: url('%(fonts)s/Archivo-Bold.ttf'); font-weight: 700; }
@font-face { font-family: Plex; src: url('%(fonts)s/IBMPlexMono-SemiBold.ttf'); }
@page { size: A4; margin: 16mm 18mm 18mm; @bottom-left { content: "%(footer)s"; font: 7pt Plex; color: %(muted)s; } }
body { margin: 0; font-family: Archivo; font-size: 9.8pt; line-height: 1.4; color: %(ink)s; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
h1 { font-size: 21pt; line-height: 1.1; margin: 0 0 4mm; }
h2 { color: %(accent)s; font-size: 12.5pt; margin: 6mm 0 2mm; break-after: avoid; }
h3 { font-size: 10.5pt; margin: 4mm 0 1.5mm; break-after: avoid; }
p { margin: 0 0 2.2mm; }
ul, ol { margin: 0 0 2.5mm; padding-left: 5mm; }
li { margin-bottom: 1mm; }
table { border-collapse: collapse; width: 100%%; margin: 1mm 0 3.5mm; font-size: 8.4pt; line-height: 1.3; }
tr { break-inside: avoid; }
th { background: %(ink)s; color: %(paper)s; font-family: Plex; font-size: 6.6pt; letter-spacing: 1pt; text-transform: uppercase; text-align: left; padding: 1.6mm 2mm; }
td { border-bottom: 0.5pt solid %(rule)s; padding: 1.4mm 2mm; vertical-align: top; }
td.id { white-space: nowrap; }
figure { margin: 2mm 0; break-inside: avoid; }
figure svg { max-width: 100%%; height: auto; display: block; margin: 0 auto; }
figcaption { margin-top: 2mm; color: %(muted)s; font-size: 8.5pt; }
figure.pic img { display: block; max-width: 140mm; max-height: 90mm; border: 0.5pt solid %(rule)s; }
figure.pic .missing { width: 140mm; height: 22mm; border: 0.75pt dashed %(muted)s; color: %(muted)s; font-size: 8pt;
                      display: flex; align-items: center; justify-content: center; }
"""


def inline(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", html.escape(t))
    t = re.sub(EMPHASIS, r"<i>\1</i>", t)
    return re.sub(r"`([^`]+)`", r"<code>\1</code>", t)


def pdf_page(blocks, theme, footer):
    parts, lst, merged, start = [], None, [], None
    for kind, val in blocks:
        if kind == "p" and merged and merged[-1][0] == "diagram" and "caption" not in merged[-1][1]:
            merged[-1] = ("diagram", dict(merged[-1][1], caption=val))
        else:
            merged.append((kind, val))
    for kind, val in merged:
        if kind == "start":
            start = val
            continue
        if kind not in ("ul", "ol") and lst:
            parts.append("</%s>" % lst)
            lst = None
        if kind in ("ul", "ol"):
            if lst != kind:
                if lst:
                    parts.append("</%s>" % lst)
                parts.append("<%s%s>" % (kind, " start='%d'" % start if start and kind == "ol" else ""))
                lst, start = kind, None
            parts.append("<li>%s</li>" % inline(val))
        elif kind == "page":
            parts.append("<div style='break-before: page'></div>")
        elif kind == "image":
            pic = "<img src='%s'>" % val["path"].as_uri() if val["path"].exists() else \
                "<div class='missing'>Picture to come: %s</div>" % inline(val["caption"])
            parts.append("<figure class='pic'>%s<figcaption>%s</figcaption></figure>" % (pic, inline(val["caption"])))
        elif kind == "diagram":
            caption = "<figcaption>%s</figcaption>" % inline(val["caption"]) if "caption" in val else ""
            parts.append("<figure>%s%s</figure>" % (val["svg"].read_text(encoding="utf-8"), caption))
        elif kind == "table":
            head = "".join("<th>%s</th>" % inline(c) for c in val[0])
            body = "".join("<tr>%s</tr>" % "".join("<td%s>%s</td>" % (" class='id'" if ID.match(c) else "", inline(c)) for c in r)
                           for r in val[1:])
            parts.append("<table><thead><tr>%s</tr></thead><tbody>%s</tbody></table>" % (head, body))
        else:
            tag = {"h1": "h1", "h2": "h2", "h3": "h3"}.get(kind, "p")
            parts.append("<%s>%s</%s>" % (tag, inline(val), tag))
    if lst:
        parts.append("</%s>" % lst)
    css = CSS % dict(theme, fonts=(TOOLS / "fonts").as_uri(), footer=footer.replace('"', "'"))
    return "<!doctype html><html><head><meta charset='utf-8'><style>%s</style></head><body>%s</body></html>" % (css, "\n".join(parts))


def pdf(blocks, out, theme, footer):
    src = out.with_suffix(".html")
    src.write_text(pdf_page(blocks, theme, footer), encoding="utf-8")
    try:
        subprocess.run(["node", str(TOOLS / "render.js"), "pdf", str(src), str(out.with_suffix(".pdf"))],
                       check=True, capture_output=True, text=True, timeout=300)
    except Exception as e:
        chrome = shutil.which("chromium") or shutil.which("chromium-browser") or shutil.which("google-chrome")
        if not chrome:
            print("  the PDF is not built (%s); the page is in %s" % (why(e), src.name))
            return False
        subprocess.run([chrome, "--headless", "--no-sandbox", "--no-pdf-header-footer",
                        "--print-to-pdf=%s" % out.with_suffix(".pdf"), src.as_uri()], check=True, capture_output=True)
    src.unlink()
    return True


def main():
    args = sys.argv[1:]
    root = TOOLS.parent
    footer = None
    if "--vault" in args:
        i = args.index("--vault")
        root = Path(args[i + 1])
        del args[i:i + 2]
    if "--footer" in args:
        i = args.index("--footer")
        footer = args[i + 1]
        del args[i:i + 2]
    if not args:
        sys.exit(__doc__)
    settings_file = Path(root) / "settings.json"
    settings = json.loads(settings_file.read_text(encoding="utf-8")) if settings_file.exists() else {}
    theme = dict(THEME, **settings.get("theme", {}))
    footer = footer if footer is not None else settings.get("footer", "")
    note = Path(args[0])
    out = Path(args[1]).resolve().with_suffix("") if len(args) > 1 else (Path(root) / "build" / note.stem).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    blocks = diagrams(parse(note), out.parent / (out.name + " figures"), theme)
    title = next((val for kind, val in blocks if kind == "h1"), note.stem)
    word(blocks, out.with_suffix(".docx"), theme, footer, title)
    made = [out.with_suffix(".docx").name]
    if pdf(blocks, out, theme, footer):
        made.append(out.with_suffix(".pdf").name)
    print("Built %s in %s" % (" and ".join(made), out.parent))


if __name__ == "__main__":
    main()
